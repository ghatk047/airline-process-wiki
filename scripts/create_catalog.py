#!/usr/bin/env python3
"""
create_catalog.py — Bootstrap Airlines_Process_Catalog.xlsx from processes.json
Creates the Index tab with all L2 rows. Run once to initialise the workbook.
Per-subprocess L4 tabs are added by the pipeline agent as each subprocess is completed.
"""
import json
import sys
from pathlib import Path
from datetime import date

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ── Paths ──────────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).parent
REPO_DIR   = SCRIPT_DIR.parent
DATA_DIR   = REPO_DIR / "data"
PROCESSES  = DATA_DIR / "processes.json"
EXCEL_OUT  = DATA_DIR / "Airlines_Process_Catalog.xlsx"

# ── Palette ────────────────────────────────────────────────────────────────
NAVY        = "003366"
ORANGE      = "FF6600"
LIGHT_BLUE  = "D6E4F0"
ALT_ROW     = "F0F7FF"
WHITE       = "FFFFFF"
STATUS_BG   = "FFF3CD"
STATUS_FG   = "7F4F00"
DONE_BG     = "D4EDDA"
DONE_FG     = "155724"

def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def header_font(bold=True, color="FFFFFF", size=10):
    return Font(name="Calibri", bold=bold, color=color, size=size)

def cell_font(bold=False, color="000000", size=9):
    return Font(name="Calibri", bold=bold, color=color, size=size)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def centre():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Load processes ─────────────────────────────────────────────────────────
with open(PROCESSES) as f:
    data = json.load(f)

processes = data["processes"]

# ── Build workbook ─────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Index"

# ── Header row 1 — title banner ────────────────────────────────────────────
ws.merge_cells("A1:G1")
title_cell = ws["A1"]
title_cell.value = "Airlines Process Catalog — Master Index"
title_cell.font  = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
title_cell.fill  = fill(NAVY)
title_cell.alignment = centre()
ws.row_dimensions[1].height = 28

# ── Header row 2 — column labels ───────────────────────────────────────────
COLS = ["Process ID", "L1 Domain", "L2 Process", "L3 Name",
        "Status", "Run Date", "Wiki Path"]
for ci, label in enumerate(COLS, 1):
    c = ws.cell(row=2, column=ci, value=label)
    c.font      = header_font(bold=True, color="FFFFFF", size=9)
    c.fill      = fill(ORANGE)
    c.alignment = centre()
    c.border    = thin_border()
ws.row_dimensions[2].height = 20

# ── Column widths ──────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 32
ws.column_dimensions["D"].width = 38
ws.column_dimensions["E"].width = 13
ws.column_dimensions["F"].width = 13
ws.column_dimensions["G"].width = 45

# ── Freeze panes ───────────────────────────────────────────────────────────
ws.freeze_panes = "A3"

# ── Data rows ──────────────────────────────────────────────────────────────
prev_l1 = ""
row_idx  = 3

for i, proc in enumerate(processes):
    pid      = proc["id"]
    l1       = proc["l1_domain"]
    l2       = proc["l2_process"]
    l3       = proc["l3_name"]
    status   = proc["status"]
    run_date = proc.get("run_date") or ""
    wiki     = proc.get("wiki_path") or ""

    # Choose row background
    is_new_l1 = (l1 != prev_l1)
    if is_new_l1:
        row_bg = LIGHT_BLUE
    elif i % 2 == 0:
        row_bg = ALT_ROW
    else:
        row_bg = WHITE

    row_data = [pid, l1, l2, l3, status, run_date, wiki]
    for ci, val in enumerate(row_data, 1):
        c = ws.cell(row=row_idx, column=ci, value=val)
        c.border    = thin_border()
        c.font      = cell_font(bold=is_new_l1 and ci <= 3, size=9)
        c.fill      = fill(row_bg)
        c.alignment = left() if ci > 1 else centre()

    # Status cell colouring
    status_cell = ws.cell(row=row_idx, column=5)
    if status == "Complete":
        status_cell.fill = fill(DONE_BG)
        status_cell.font = cell_font(bold=True, color=DONE_FG, size=9)
        status_cell.value = "✅ Complete"
    elif status == "In Progress":
        status_cell.fill = fill("D0ECF7")
        status_cell.font = cell_font(bold=True, color="0C5D8A", size=9)
        status_cell.value = "🔄 In Progress"
    else:
        status_cell.fill = fill(STATUS_BG)
        status_cell.font = cell_font(color=STATUS_FG, size=9)
        status_cell.value = "⏳ Queued"
    status_cell.alignment = centre()

    ws.row_dimensions[row_idx].height = 16
    prev_l1 = l1
    row_idx += 1

# ── Summary row ────────────────────────────────────────────────────────────
total       = len(processes)
complete    = sum(1 for p in processes if p["status"] == "Complete")
in_progress = sum(1 for p in processes if p["status"] == "In Progress")
queued      = total - complete - in_progress

ws.merge_cells(f"A{row_idx}:D{row_idx}")
summary_cell = ws[f"A{row_idx}"]
summary_cell.value = (
    f"Total: {total} processes  |  "
    f"✅ Complete: {complete}  |  "
    f"🔄 In Progress: {in_progress}  |  "
    f"⏳ Queued: {queued}"
)
summary_cell.font      = cell_font(bold=True, color="FFFFFF", size=9)
summary_cell.fill      = fill(NAVY)
summary_cell.alignment = left()
ws.row_dimensions[row_idx].height = 18

# ── Save ───────────────────────────────────────────────────────────────────
EXCEL_OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(EXCEL_OUT)
print(f"✅ Created: {EXCEL_OUT}")
print(f"   Rows: {total}  |  Complete: {complete}  |  Queued: {queued}")
