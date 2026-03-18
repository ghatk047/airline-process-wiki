#!/usr/bin/env python3
"""
write_excel.py — Thread/process-safe Excel writer for pipeline agents.
Each per-L2 agent calls this to:
  1. Create a dedicated subprocess tab with full L4 rows
  2. Append those rows to the Master Catalog tab
  3. Update the Index tab status + run date

Usage:
    python3 write_excel.py --process-id NP-SP-03 --json-file /path/to/np-sp-03.json

The JSON file must match the per-subprocess schema in processes.json.
File locking via fcntl (Unix) prevents corruption when agents run in parallel.
"""
import argparse
import fcntl
import json
import sys
from datetime import date
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ── Paths ──────────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).parent
REPO_DIR   = SCRIPT_DIR.parent
DATA_DIR   = REPO_DIR / "data"
EXCEL_FILE = DATA_DIR / "Airlines_Process_Catalog.xlsx"
LOCK_FILE  = DATA_DIR / ".excel.lock"

# ── L4 columns ─────────────────────────────────────────────────────────────
L4_COLS = [
    "Process ID", "L1 Domain", "L2 Process", "L3 Name",
    "L4 Step #", "L4 Step Name", "Role / Swim Lane", "System",
    "Input", "Output", "KPI", "Pain Point / Risk",
    "Decision Point", "Exception"
]

# ── Domain colour map (header fill per L1) ─────────────────────────────────
DOMAIN_COLOURS = {
    "Network Planning & Scheduling": "003366",
    "Customer Experience & Loyalty": "5B2D8E",
    "Flight Operations":             "1A5276",
    "Crew Management":               "154360",
    "Ground Operations & Airport Services": "145A32",
    "Maintenance, Repair & Overhaul":       "784212",
    "Corporate Support Functions":          "4A235A",
}
DEFAULT_COLOUR = "2C3E50"

# ── Style helpers ──────────────────────────────────────────────────────────
def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def hfont(size=9):
    return Font(name="Calibri", bold=True, color="FFFFFF", size=size)

def dfont(bold=False, size=9):
    return Font(name="Calibri", bold=bold, color="000000", size=size)

def centre():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Core writer ────────────────────────────────────────────────────────────
def write_subprocess_tab(wb, proc: dict):
    """Add or replace a tab named after the process ID with full L4 rows."""
    pid   = proc["id"]
    l1    = proc["l1_domain"]
    l2    = proc["l2_process"]
    l3    = proc["l3_name"]
    steps = proc.get("l4_steps", [])

    # Remove existing tab if any
    if pid in wb.sheetnames:
        del wb[pid]

    ws = wb.create_sheet(title=pid)

    # Header colour
    hdr_colour = DOMAIN_COLOURS.get(l1, DEFAULT_COLOUR)

    # Row 1 — title banner
    ws.merge_cells(f"A1:{get_column_letter(len(L4_COLS))}1")
    tc = ws["A1"]
    tc.value     = f"{pid} — {l3}"
    tc.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    tc.fill      = fill(hdr_colour)
    tc.alignment = centre()
    ws.row_dimensions[1].height = 26

    # Row 2 — subtitle
    ws.merge_cells(f"A2:{get_column_letter(len(L4_COLS))}2")
    sc = ws["A2"]
    sc.value     = f"L1: {l1}  |  L2: {l2}  |  L3: {l3}"
    sc.font      = Font(name="Calibri", bold=False, color="FFFFFF", size=9)
    sc.fill      = fill(hdr_colour)
    sc.alignment = centre()
    ws.row_dimensions[2].height = 16

    # Row 3 — column headers
    for ci, col in enumerate(L4_COLS, 1):
        c = ws.cell(row=3, column=ci, value=col)
        c.font      = hfont()
        c.fill      = fill("FF6600")
        c.alignment = centre()
        c.border    = thin_border()
    ws.row_dimensions[3].height = 20

    # Freeze pane
    ws.freeze_panes = "A4"

    # Column widths
    WIDTHS = [11, 28, 26, 28, 9, 36, 24, 28, 28, 28, 26, 36, 14, 12]
    for ci, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # Data rows
    for ri, step in enumerate(steps, 4):
        row_bg = "F0F7FF" if ri % 2 == 0 else "FFFFFF"
        row_vals = [
            pid,
            l1,
            l2,
            l3,
            step.get("step", ""),
            step.get("name", ""),
            step.get("role", ""),
            step.get("system", ""),
            step.get("input", ""),
            step.get("output", ""),
            step.get("kpi", ""),
            step.get("pain_point", ""),
            step.get("decision_point", "N"),
            step.get("exception", "N"),
        ]
        for ci, val in enumerate(row_vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font      = dfont(size=8)
            c.fill      = fill(row_bg)
            c.alignment = left() if ci not in (1, 5, 13, 14) else centre()
            c.border    = thin_border()
        ws.row_dimensions[ri].height = 30

    print(f"  ↳ Tab '{pid}' written ({len(steps)} L4 steps)")
    return len(steps)


def append_to_master(wb, proc: dict):
    """Append L4 rows to Master Catalog tab (create tab if needed)."""
    steps = proc.get("l4_steps", [])
    pid   = proc["id"]
    l1    = proc["l1_domain"]
    l2    = proc["l2_process"]
    l3    = proc["l3_name"]

    TAB = "Master Catalog"
    if TAB not in wb.sheetnames:
        ws = wb.create_sheet(title=TAB, index=1)
        # Header
        ws.merge_cells(f"A1:{get_column_letter(len(L4_COLS))}1")
        tc = ws["A1"]
        tc.value     = "Airlines Process Catalog — Master L4 Row Catalog"
        tc.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
        tc.fill      = fill("003366")
        tc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26
        for ci, col in enumerate(L4_COLS, 1):
            c = ws.cell(row=2, column=ci, value=col)
            c.font      = hfont()
            c.fill      = fill("FF6600")
            c.alignment = centre()
            c.border    = thin_border()
        ws.row_dimensions[2].height = 20
        ws.freeze_panes = "A3"
        WIDTHS = [11, 28, 26, 28, 9, 36, 24, 28, 28, 28, 26, 36, 14, 12]
        for ci, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        next_row = 3
    else:
        ws = wb[TAB]
        next_row = ws.max_row + 1

    hdr_colour = DOMAIN_COLOURS.get(l1, DEFAULT_COLOUR)
    for step in steps:
        row_bg = "F0F7FF" if next_row % 2 == 0 else "FFFFFF"
        row_vals = [
            pid, l1, l2, l3,
            step.get("step", ""),
            step.get("name", ""),
            step.get("role", ""),
            step.get("system", ""),
            step.get("input", ""),
            step.get("output", ""),
            step.get("kpi", ""),
            step.get("pain_point", ""),
            step.get("decision_point", "N"),
            step.get("exception", "N"),
        ]
        for ci, val in enumerate(row_vals, 1):
            c = ws.cell(row=next_row, column=ci, value=val)
            c.font      = dfont(size=8)
            c.fill      = fill(row_bg)
            c.alignment = left() if ci not in (1, 5, 13, 14) else centre()
            c.border    = thin_border()
        ws.row_dimensions[next_row].height = 30
        next_row += 1

    print(f"  ↳ Master Catalog: {len(steps)} rows appended")


def update_index(wb, proc: dict):
    """Set Status = Complete and Run Date on the Index tab."""
    if "Index" not in wb.sheetnames:
        print("  ⚠ Index tab not found — skipping status update")
        return

    ws     = wb["Index"]
    pid    = proc["id"]
    today  = date.today().isoformat()

    for row in ws.iter_rows(min_row=3):
        if row[0].value == pid:
            status_cell   = row[4]
            rundate_cell  = row[5]
            wiki_cell     = row[6]
            status_cell.value     = "✅ Complete"
            status_cell.font      = Font(name="Calibri", bold=True, color="155724", size=9)
            status_cell.fill      = PatternFill("solid", fgColor="D4EDDA")
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            rundate_cell.value    = today
            if proc.get("wiki_path"):
                wiki_cell.value = proc["wiki_path"]
            print(f"  ↳ Index updated: {pid} → Complete ({today})")
            return
    print(f"  ⚠ Index row for {pid} not found")


# ── Main ───────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Write subprocess data to Excel catalog")
    parser.add_argument("--process-id", required=True, help="e.g. NP-SP-03")
    parser.add_argument("--json-file",  required=True, help="Path to per-subprocess JSON")
    parser.add_argument("--excel-file", default=str(EXCEL_FILE), help="Path to Excel workbook")
    args = parser.parse_args()

    json_path  = Path(args.json_file)
    excel_path = Path(args.excel_file)

    if not json_path.exists():
        print(f"ERROR: JSON file not found: {json_path}")
        sys.exit(1)
    if not excel_path.exists():
        print(f"ERROR: Excel file not found: {excel_path}")
        print("Run create_catalog.py first to bootstrap the workbook.")
        sys.exit(1)

    with open(json_path) as f:
        proc = json.load(f)

    print(f"\n📊 Writing {args.process_id} to Excel...")

    # File lock — prevents corruption if agents run in parallel
    lock_path = excel_path.parent / ".excel.lock"
    with open(lock_path, "w") as lock_fh:
        fcntl.flock(lock_fh, fcntl.LOCK_EX)
        try:
            wb = openpyxl.load_workbook(excel_path)
            write_subprocess_tab(wb, proc)
            append_to_master(wb, proc)
            update_index(wb, proc)
            wb.save(excel_path)
            print(f"✅ Saved: {excel_path}")
        finally:
            fcntl.flock(lock_fh, fcntl.LOCK_UN)


if __name__ == "__main__":
    main()
